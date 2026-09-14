"""Reporting.

Every report returns the same envelope — `columns`, `rows`, `summary`, `charts`
— so one screen renders all of them and adding a report is a function, not a
new page. Exports reuse the exact rows the screen showed, which is the only way
a printed report and the screen can be guaranteed to agree.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from typing import Any

from bson import ObjectId

from app.core.context import TenantContext
from app.core.exceptions import NotFound
from app.db.mongo import C, collection
from app.models.base import to_datetime

PRESENT_LIKE = ["present", "late", "half_day"]


@dataclass
class Column:
    key: str
    label: str
    type: str = "text"           # text | number | money | percent | date | badge
    align: str = "left"


@dataclass
class Report:
    key: str
    title: str
    description: str
    columns: list[Column]
    rows: list[dict[str, Any]]
    summary: list[dict[str, Any]] = field(default_factory=list)
    charts: list[dict[str, Any]] = field(default_factory=list)
    generated_at: datetime = field(default_factory=lambda: datetime.now(UTC))

    def as_dict(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "title": self.title,
            "description": self.description,
            "columns": [c.__dict__ for c in self.columns],
            "rows": self.rows,
            "summary": self.summary,
            "charts": self.charts,
            "row_count": len(self.rows),
            "generated_at": self.generated_at.isoformat(),
        }


def _range(start: date | None, end: date | None, default_days: int = 30) -> tuple[date, date]:
    end = end or date.today()
    start = start or (end - timedelta(days=default_days))
    return start, end


def _between(start: date, end: date) -> dict[str, Any]:
    return {"$gte": to_datetime(start), "$lte": to_datetime(end) + timedelta(days=1)}


async def _lookup(name: str, tenant_id: ObjectId, field_name: str = "name") -> dict[ObjectId, str]:
    rows = await collection(name).find(
        {"tenant_id": tenant_id, "is_deleted": {"$ne": True}}, {field_name: 1}
    ).to_list(length=None)
    return {row["_id"]: row.get(field_name, "") for row in rows}


async def _subject_labels(tenant_id: ObjectId) -> dict[ObjectId, str]:
    """Subjects repeat across classes — "Mathematics" exists for Class 7 and
    Class 8 — so a bare name makes two different rows look like duplicates."""
    subjects = await collection(C.SUBJECTS).find(
        {"tenant_id": tenant_id, "is_deleted": {"$ne": True}}, {"name": 1, "class_id": 1}
    ).to_list(length=None)
    classes = await _lookup(C.CLASSES, tenant_id)
    names: dict[str, int] = {}
    for subject in subjects:
        names[subject.get("name", "")] = names.get(subject.get("name", ""), 0) + 1
    return {
        subject["_id"]: (
            f"{subject.get('name', '')} · {classes.get(subject.get('class_id'), '')}".strip(" ·")
            if names.get(subject.get("name", ""), 0) > 1
            else subject.get("name", "")
        )
        for subject in subjects
    }


# ── Catalogue ─────────────────────────────────────────────────────────────
REPORT_CATALOGUE = [
    {"key": "student_roster", "title": "Student Roster", "module": "students",
     "group": "People", "description": "Every enrolled student with class, contact and dues.",
     "icon": "users"},
    {"key": "student_demographics", "title": "Student Demographics", "module": "students",
     "group": "People", "description": "Gender, category and class distribution.",
     "icon": "chart-pie"},
    {"key": "staff_directory", "title": "Staff Directory", "module": "staff",
     "group": "People", "description": "Staff by department, designation and status.",
     "icon": "briefcase"},
    {"key": "attendance_register", "title": "Attendance Register", "module": "attendance",
     "group": "Attendance", "description": "Per-student attendance over a date range.",
     "icon": "check-square"},
    {"key": "attendance_defaulters", "title": "Attendance Defaulters", "module": "attendance",
     "group": "Attendance", "description": "Students below an attendance threshold.",
     "icon": "triangle-alert"},
    {"key": "daily_attendance", "title": "Daily Attendance Trend", "module": "attendance",
     "group": "Attendance", "description": "Section-wise attendance day by day.",
     "icon": "calendar-days"},
    {"key": "fee_collection", "title": "Fee Collection", "module": "payments",
     "group": "Finance", "description": "Receipts collected, by day and by method.",
     "icon": "wallet"},
    {"key": "fee_outstanding", "title": "Outstanding Dues", "module": "invoices",
     "group": "Finance", "description": "Who owes what, oldest first.",
     "icon": "receipt"},
    {"key": "fee_by_class", "title": "Fees by Class", "module": "invoices",
     "group": "Finance", "description": "Billed against collected, per class.",
     "icon": "layers"},
    {"key": "expense_summary", "title": "Expense Summary", "module": "expenses",
     "group": "Finance", "description": "Spending by category over a period.",
     "icon": "trending-down"},
    {"key": "exam_performance", "title": "Exam Performance", "module": "exams",
     "group": "Academics", "description": "Subject averages and pass rates for an exam.",
     "icon": "file-badge"},
    {"key": "student_marksheet", "title": "Class Marksheet", "module": "exams",
     "group": "Academics", "description": "Every student's marks across subjects, ranked.",
     "icon": "trophy"},
    {"key": "admission_funnel", "title": "Admission Funnel", "module": "admissions",
     "group": "Admissions", "description": "Enquiries through to enrolment, by source.",
     "icon": "user-plus"},
    {"key": "library_circulation", "title": "Library Circulation", "module": "library",
     "group": "Facilities", "description": "Issued, returned and overdue items.",
     "icon": "book-marked"},
    {"key": "transport_usage", "title": "Transport Usage", "module": "transport",
     "group": "Facilities", "description": "Route occupancy and fare collection.",
     "icon": "bus"},
]


def catalogue_for(tenant: TenantContext, permissions: set[str]) -> list[dict[str, Any]]:
    from app.core.permissions import has_permission

    return [
        report for report in REPORT_CATALOGUE
        if tenant.module_enabled(report["module"])
        and has_permission(permissions, f"{report['module']}:read")
    ]


# ── Reports ───────────────────────────────────────────────────────────────
async def student_roster(tenant: TenantContext, **params) -> Report:
    query: dict[str, Any] = {"tenant_id": tenant.id, "is_deleted": {"$ne": True}}
    if params.get("status"):
        query["status"] = params["status"]
    else:
        query["status"] = "active"
    if params.get("class_id"):
        query["current_class_id"] = ObjectId(params["class_id"])
    if params.get("section_id"):
        query["current_section_id"] = ObjectId(params["section_id"])

    students = await collection(C.STUDENTS).find(query).sort(
        [("current_class_id", 1), ("roll_number", 1)]
    ).to_list(length=5000)

    classes = await _lookup(C.CLASSES, tenant.id)
    sections = await _lookup(C.SECTIONS, tenant.id)

    rows = [
        {
            "admission_number": s.get("admission_number", ""),
            "roll_number": s.get("roll_number", ""),
            "name": " ".join(filter(None, [s.get("first_name"), s.get("last_name")])),
            "class": classes.get(s.get("current_class_id"), "—"),
            "section": sections.get(s.get("current_section_id"), "—"),
            "gender": (s.get("gender") or "").title(),
            "phone": (s.get("contact") or {}).get("phone", ""),
            "outstanding": round(float(s.get("outstanding_amount") or 0), 2),
            "status": s.get("status", ""),
        }
        for s in students
    ]

    dues = sum(r["outstanding"] for r in rows)
    return Report(
        key="student_roster",
        title="Student Roster",
        description="Every enrolled student with class, contact and dues.",
        columns=[
            Column("admission_number", "Admission No."),
            Column("roll_number", "Roll", "number", "center"),
            Column("name", "Student"),
            Column("class", "Class"),
            Column("section", "Section"),
            Column("gender", "Gender"),
            Column("phone", "Phone"),
            Column("outstanding", "Dues", "money", "right"),
            Column("status", "Status", "badge", "center"),
        ],
        rows=rows,
        summary=[
            {"label": "Students", "value": len(rows)},
            {"label": "Total dues", "value": round(dues, 2), "type": "money"},
            {"label": "With dues", "value": sum(1 for r in rows if r["outstanding"] > 0)},
        ],
    )


async def student_demographics(tenant: TenantContext, **params) -> Report:
    base = {"tenant_id": tenant.id, "status": "active", "is_deleted": {"$ne": True}}

    by_class = await collection(C.STUDENTS).aggregate([
        {"$match": base},
        {"$group": {
            "_id": "$current_class_id",
            "total": {"$sum": 1},
            "female": {"$sum": {"$cond": [{"$eq": ["$gender", "female"]}, 1, 0]}},
            "male": {"$sum": {"$cond": [{"$eq": ["$gender", "male"]}, 1, 0]}},
            "transport": {"$sum": {"$cond": ["$uses_transport", 1, 0]}},
            "hostel": {"$sum": {"$cond": ["$is_hosteller", 1, 0]}},
        }},
        {"$sort": {"total": -1}},
    ]).to_list(length=None)

    classes = await _lookup(C.CLASSES, tenant.id)
    rows = [
        {
            "class": classes.get(row["_id"], "Unassigned"),
            "total": row["total"],
            "female": row["female"],
            "male": row["male"],
            "transport": row["transport"],
            "hostel": row["hostel"],
        }
        for row in by_class
    ]

    by_gender = await collection(C.STUDENTS).aggregate([
        {"$match": base}, {"$group": {"_id": "$gender", "n": {"$sum": 1}}},
    ]).to_list(length=None)
    by_category = await collection(C.STUDENTS).aggregate([
        {"$match": base}, {"$group": {"_id": "$category", "n": {"$sum": 1}}},
        {"$sort": {"n": -1}}, {"$limit": 8},
    ]).to_list(length=None)

    return Report(
        key="student_demographics",
        title="Student Demographics",
        description="Gender, category and class distribution.",
        columns=[
            Column("class", "Class"),
            Column("total", "Students", "number", "right"),
            Column("female", "Female", "number", "right"),
            Column("male", "Male", "number", "right"),
            Column("transport", "Transport", "number", "right"),
            Column("hostel", "Hostel", "number", "right"),
        ],
        rows=rows,
        summary=[
            {"label": "Total students", "value": sum(r["total"] for r in rows)},
            {"label": "Using transport", "value": sum(r["transport"] for r in rows)},
            {"label": "Hostellers", "value": sum(r["hostel"] for r in rows)},
        ],
        charts=[
            {"type": "donut", "title": "By gender",
             "data": [{"name": (r["_id"] or "undisclosed").title(), "value": r["n"]}
                      for r in by_gender]},
            {"type": "bar", "title": "By category",
             "data": [{"name": r["_id"] or "Unspecified", "value": r["n"]}
                      for r in by_category]},
            {"type": "bar", "title": "By class",
             "data": [{"name": r["class"], "value": r["total"]} for r in rows[:12]]},
        ],
    )


async def staff_directory(tenant: TenantContext, **params) -> Report:
    query: dict[str, Any] = {"tenant_id": tenant.id, "is_deleted": {"$ne": True}}
    query["status"] = params.get("status") or "active"
    if params.get("department_id"):
        query["department_id"] = ObjectId(params["department_id"])

    staff = await collection(C.STAFF).find(query).sort([("first_name", 1)]).to_list(length=2000)
    departments = await _lookup(C.DEPARTMENTS, tenant.id)

    rows = [
        {
            "employee_id": s.get("employee_id", ""),
            "name": " ".join(filter(None, [s.get("first_name"), s.get("last_name")])),
            "designation": s.get("designation", ""),
            "department": departments.get(s.get("department_id"), "—"),
            "type": (s.get("employment_type") or "").replace("_", " ").title(),
            "teaching": "Yes" if s.get("is_teaching") else "No",
            "phone": (s.get("contact") or {}).get("phone", ""),
            "status": s.get("status", ""),
        }
        for s in staff
    ]
    return Report(
        key="staff_directory",
        title="Staff Directory",
        description="Staff by department, designation and status.",
        columns=[
            Column("employee_id", "Employee ID"),
            Column("name", "Name"),
            Column("designation", "Designation"),
            Column("department", "Department"),
            Column("type", "Type"),
            Column("teaching", "Teaching", "text", "center"),
            Column("phone", "Phone"),
            Column("status", "Status", "badge", "center"),
        ],
        rows=rows,
        summary=[
            {"label": "Staff", "value": len(rows)},
            {"label": "Teaching", "value": sum(1 for r in rows if r["teaching"] == "Yes")},
            {"label": "Non-teaching", "value": sum(1 for r in rows if r["teaching"] == "No")},
        ],
    )


async def attendance_register(tenant: TenantContext, **params) -> Report:
    start, end = _range(params.get("start"), params.get("end"))
    match: dict[str, Any] = {
        "tenant_id": tenant.id, "session_key": "day", "is_deleted": {"$ne": True},
        "date": _between(start, end),
    }
    if params.get("section_id"):
        match["section_id"] = ObjectId(params["section_id"])
    if params.get("class_id"):
        match["class_id"] = ObjectId(params["class_id"])

    grouped = await collection(C.ATTENDANCE).aggregate([
        {"$match": match},
        {"$group": {
            "_id": "$student_id",
            "total": {"$sum": 1},
            "present": {"$sum": {"$cond": [{"$eq": ["$status", "present"]}, 1, 0]}},
            "absent": {"$sum": {"$cond": [{"$eq": ["$status", "absent"]}, 1, 0]}},
            "late": {"$sum": {"$cond": [{"$eq": ["$status", "late"]}, 1, 0]}},
            "leave": {"$sum": {"$cond": [{"$eq": ["$status", "leave"]}, 1, 0]}},
            "counted": {"$sum": {"$cond": [{"$in": ["$status", PRESENT_LIKE]}, 1, 0]}},
        }},
    ]).to_list(length=5000)

    students = {
        s["_id"]: s
        for s in await collection(C.STUDENTS).find(
            {"_id": {"$in": [g["_id"] for g in grouped]}},
            {"first_name": 1, "last_name": 1, "roll_number": 1, "admission_number": 1,
             "current_class_id": 1, "current_section_id": 1},
        ).to_list(length=None)
    }
    classes = await _lookup(C.CLASSES, tenant.id)
    sections = await _lookup(C.SECTIONS, tenant.id)

    rows = []
    for group in grouped:
        student = students.get(group["_id"], {})
        percentage = round(group["counted"] / group["total"] * 100, 2) if group["total"] else 0
        rows.append({
            "roll_number": student.get("roll_number", ""),
            "admission_number": student.get("admission_number", ""),
            "name": " ".join(filter(None, [student.get("first_name"), student.get("last_name")])),
            "class": classes.get(student.get("current_class_id"), "—"),
            "section": sections.get(student.get("current_section_id"), "—"),
            "days": group["total"],
            "present": group["present"],
            "absent": group["absent"],
            "late": group["late"],
            "leave": group["leave"],
            "percentage": percentage,
        })
    rows.sort(key=lambda r: r["percentage"])

    overall = (
        round(sum(r["present"] + r["late"] for r in rows) /
              max(sum(r["days"] for r in rows), 1) * 100, 2)
    )
    return Report(
        key="attendance_register",
        title="Attendance Register",
        description=f"{start:%d %b %Y} to {end:%d %b %Y}",
        columns=[
            Column("roll_number", "Roll", "number", "center"),
            Column("name", "Student"),
            Column("class", "Class"),
            Column("section", "Section"),
            Column("days", "Days", "number", "right"),
            Column("present", "Present", "number", "right"),
            Column("absent", "Absent", "number", "right"),
            Column("late", "Late", "number", "right"),
            Column("percentage", "Attendance", "percent", "right"),
        ],
        rows=rows,
        summary=[
            {"label": "Students", "value": len(rows)},
            {"label": "Overall attendance", "value": overall, "type": "percent"},
            {"label": "Below 75%", "value": sum(1 for r in rows if r["percentage"] < 75)},
        ],
    )


async def attendance_defaulters(tenant: TenantContext, **params) -> Report:
    threshold = float(params.get("threshold") or 75)
    report = await attendance_register(tenant, **params)
    rows = [r for r in report.rows if r["percentage"] < threshold]
    return Report(
        key="attendance_defaulters",
        title="Attendance Defaulters",
        description=f"Below {threshold:g}% · {report.description}",
        columns=report.columns,
        rows=rows,
        summary=[
            {"label": "Defaulters", "value": len(rows)},
            {"label": "Threshold", "value": threshold, "type": "percent"},
            {"label": "Lowest", "value": rows[0]["percentage"] if rows else 0, "type": "percent"},
        ],
    )


async def daily_attendance(tenant: TenantContext, **params) -> Report:
    start, end = _range(params.get("start"), params.get("end"))
    match: dict[str, Any] = {
        "tenant_id": tenant.id, "session_key": "day", "is_deleted": {"$ne": True},
        "date": _between(start, end),
    }
    if params.get("section_id"):
        match["section_id"] = ObjectId(params["section_id"])

    grouped = await collection(C.ATTENDANCE).aggregate([
        {"$match": match},
        {"$group": {
            "_id": {"$dateToString": {"format": "%Y-%m-%d", "date": "$date"}},
            "total": {"$sum": 1},
            "present": {"$sum": {"$cond": [{"$in": ["$status", PRESENT_LIKE]}, 1, 0]}},
            "absent": {"$sum": {"$cond": [{"$eq": ["$status", "absent"]}, 1, 0]}},
        }},
        {"$sort": {"_id": 1}},
    ]).to_list(length=None)

    rows = [
        {
            "date": row["_id"],
            "total": row["total"],
            "present": row["present"],
            "absent": row["absent"],
            "percentage": round(row["present"] / row["total"] * 100, 2) if row["total"] else 0,
        }
        for row in grouped
    ]
    return Report(
        key="daily_attendance",
        title="Daily Attendance Trend",
        description=f"{start:%d %b %Y} to {end:%d %b %Y}",
        columns=[
            Column("date", "Date", "date"),
            Column("total", "Marked", "number", "right"),
            Column("present", "Present", "number", "right"),
            Column("absent", "Absent", "number", "right"),
            Column("percentage", "Rate", "percent", "right"),
        ],
        rows=rows,
        summary=[
            {"label": "Days", "value": len(rows)},
            {"label": "Average",
             "value": round(sum(r["percentage"] for r in rows) / max(len(rows), 1), 2),
             "type": "percent"},
        ],
        charts=[{"type": "line", "title": "Attendance rate",
                 "data": [{"name": r["date"], "value": r["percentage"]} for r in rows]}],
    )


async def fee_collection(tenant: TenantContext, **params) -> Report:
    start, end = _range(params.get("start"), params.get("end"))
    match: dict[str, Any] = {
        "tenant_id": tenant.id, "status": "success", "is_deleted": {"$ne": True},
        "paid_at": _between(start, end),
    }
    payments = await collection(C.PAYMENTS).find(match).sort(
        [("paid_at", -1)]
    ).to_list(length=5000)

    students = {
        s["_id"]: s
        for s in await collection(C.STUDENTS).find(
            {"_id": {"$in": [p["student_id"] for p in payments]}},
            {"first_name": 1, "last_name": 1, "admission_number": 1, "current_class_id": 1},
        ).to_list(length=None)
    }
    classes = await _lookup(C.CLASSES, tenant.id)

    rows = []
    for payment in payments:
        student = students.get(payment.get("student_id"), {})
        rows.append({
            "receipt_number": payment.get("receipt_number", ""),
            "date": payment["paid_at"].date().isoformat() if payment.get("paid_at") else "",
            "student": " ".join(filter(None, [student.get("first_name"),
                                              student.get("last_name")])),
            "admission_number": student.get("admission_number", ""),
            "class": classes.get(student.get("current_class_id"), "—"),
            "invoice": payment.get("invoice_number", ""),
            "method": (payment.get("method") or "").replace("_", " ").title(),
            "amount": round(float(payment.get("amount") or 0), 2),
        })

    by_method: dict[str, float] = {}
    by_day: dict[str, float] = {}
    for row in rows:
        by_method[row["method"]] = by_method.get(row["method"], 0) + row["amount"]
        by_day[row["date"]] = by_day.get(row["date"], 0) + row["amount"]

    total = round(sum(r["amount"] for r in rows), 2)
    return Report(
        key="fee_collection",
        title="Fee Collection",
        description=f"{start:%d %b %Y} to {end:%d %b %Y}",
        columns=[
            Column("receipt_number", "Receipt"),
            Column("date", "Date", "date"),
            Column("student", "Student"),
            Column("class", "Class"),
            Column("invoice", "Invoice"),
            Column("method", "Method", "badge"),
            Column("amount", "Amount", "money", "right"),
        ],
        rows=rows,
        summary=[
            {"label": "Collected", "value": total, "type": "money"},
            {"label": "Receipts", "value": len(rows)},
            {"label": "Average",
             "value": round(total / max(len(rows), 1), 2), "type": "money"},
        ],
        charts=[
            {"type": "donut", "title": "By method",
             "data": [{"name": k, "value": round(v, 2)} for k, v in by_method.items()]},
            {"type": "bar", "title": "By day",
             "data": [{"name": k, "value": round(v, 2)} for k, v in sorted(by_day.items())]},
        ],
    )


async def fee_outstanding(tenant: TenantContext, **params) -> Report:
    query: dict[str, Any] = {
        "tenant_id": tenant.id, "is_deleted": {"$ne": True},
        "status": {"$in": ["issued", "partially_paid", "overdue"]},
    }
    if params.get("class_id"):
        query["class_id"] = ObjectId(params["class_id"])

    invoices = await collection(C.FEE_INVOICES).find(query).sort(
        [("due_date", 1)]
    ).to_list(length=5000)

    students = {
        s["_id"]: s
        for s in await collection(C.STUDENTS).find(
            {"_id": {"$in": [i["student_id"] for i in invoices]}},
            {"first_name": 1, "last_name": 1, "admission_number": 1, "contact": 1,
             "current_class_id": 1},
        ).to_list(length=None)
    }
    classes = await _lookup(C.CLASSES, tenant.id)
    today = date.today()

    rows = []
    for invoice in invoices:
        student = students.get(invoice.get("student_id"), {})
        balance = round(float(invoice.get("total", 0)) - float(invoice.get("paid_amount", 0)), 2)
        if balance <= 0:
            continue
        due = invoice.get("due_date")
        overdue_days = (today - due.date()).days if due else 0
        rows.append({
            "invoice": invoice.get("number", ""),
            "student": " ".join(filter(None, [student.get("first_name"),
                                              student.get("last_name")])),
            "admission_number": student.get("admission_number", ""),
            "class": classes.get(student.get("current_class_id"), "—"),
            "phone": (student.get("contact") or {}).get("phone", ""),
            "period": invoice.get("period_label", ""),
            "due_date": due.date().isoformat() if due else "",
            "overdue_days": max(overdue_days, 0),
            "total": round(float(invoice.get("total") or 0), 2),
            "paid": round(float(invoice.get("paid_amount") or 0), 2),
            "balance": balance,
            "status": invoice.get("status", ""),
        })

    total_due = round(sum(r["balance"] for r in rows), 2)
    return Report(
        key="fee_outstanding",
        title="Outstanding Dues",
        description="Unpaid and part-paid invoices, oldest first.",
        columns=[
            Column("invoice", "Invoice"),
            Column("student", "Student"),
            Column("class", "Class"),
            Column("phone", "Phone"),
            Column("period", "Period"),
            Column("due_date", "Due", "date"),
            Column("overdue_days", "Days late", "number", "right"),
            Column("balance", "Balance", "money", "right"),
            Column("status", "Status", "badge", "center"),
        ],
        rows=rows,
        summary=[
            {"label": "Outstanding", "value": total_due, "type": "money"},
            {"label": "Invoices", "value": len(rows)},
            {"label": "Overdue", "value": sum(1 for r in rows if r["overdue_days"] > 0)},
        ],
    )


async def fee_by_class(tenant: TenantContext, **params) -> Report:
    grouped = await collection(C.FEE_INVOICES).aggregate([
        {"$match": {"tenant_id": tenant.id, "is_deleted": {"$ne": True},
                    "status": {"$nin": ["draft", "cancelled"]}}},
        {"$group": {
            "_id": "$class_id",
            "billed": {"$sum": "$total"},
            "collected": {"$sum": "$paid_amount"},
            "invoices": {"$sum": 1},
        }},
        {"$sort": {"billed": -1}},
    ]).to_list(length=None)

    classes = await _lookup(C.CLASSES, tenant.id)
    rows = []
    for group in grouped:
        billed = round(float(group["billed"] or 0), 2)
        collected = round(float(group["collected"] or 0), 2)
        rows.append({
            "class": classes.get(group["_id"], "Unassigned"),
            "invoices": group["invoices"],
            "billed": billed,
            "collected": collected,
            "outstanding": round(billed - collected, 2),
            "rate": round(collected / billed * 100, 2) if billed else 0,
        })

    total_billed = round(sum(r["billed"] for r in rows), 2)
    total_collected = round(sum(r["collected"] for r in rows), 2)
    return Report(
        key="fee_by_class",
        title="Fees by Class",
        description="Billed against collected, per class.",
        columns=[
            Column("class", "Class"),
            Column("invoices", "Invoices", "number", "right"),
            Column("billed", "Billed", "money", "right"),
            Column("collected", "Collected", "money", "right"),
            Column("outstanding", "Outstanding", "money", "right"),
            Column("rate", "Collection", "percent", "right"),
        ],
        rows=rows,
        summary=[
            {"label": "Billed", "value": total_billed, "type": "money"},
            {"label": "Collected", "value": total_collected, "type": "money"},
            {"label": "Collection rate",
             "value": round(total_collected / total_billed * 100, 2) if total_billed else 0,
             "type": "percent"},
        ],
        charts=[{"type": "grouped-bar", "title": "Billed vs collected",
                 "data": [{"name": r["class"], "billed": r["billed"],
                           "collected": r["collected"]} for r in rows[:12]]}],
    )


async def expense_summary(tenant: TenantContext, **params) -> Report:
    start, end = _range(params.get("start"), params.get("end"), 90)
    grouped = await collection(C.EXPENSES).aggregate([
        {"$match": {"tenant_id": tenant.id, "is_deleted": {"$ne": True},
                    "spent_on": _between(start, end)}},
        {"$group": {"_id": "$category", "total": {"$sum": "$amount"},
                    "count": {"$sum": 1}}},
        {"$sort": {"total": -1}},
    ]).to_list(length=None)

    rows = [
        {
            "category": (group["_id"] or "general").replace("_", " ").title(),
            "count": group["count"],
            "total": round(float(group["total"] or 0), 2),
        }
        for group in grouped
    ]
    total = round(sum(r["total"] for r in rows), 2)
    for row in rows:
        row["share"] = round(row["total"] / total * 100, 2) if total else 0

    return Report(
        key="expense_summary",
        title="Expense Summary",
        description=f"{start:%d %b %Y} to {end:%d %b %Y}",
        columns=[
            Column("category", "Category"),
            Column("count", "Entries", "number", "right"),
            Column("total", "Amount", "money", "right"),
            Column("share", "Share", "percent", "right"),
        ],
        rows=rows,
        summary=[{"label": "Total spend", "value": total, "type": "money"},
                 {"label": "Entries", "value": sum(r["count"] for r in rows)}],
        charts=[{"type": "donut", "title": "By category",
                 "data": [{"name": r["category"], "value": r["total"]} for r in rows]}],
    )


async def exam_performance(tenant: TenantContext, **params) -> Report:
    exam_id = params.get("exam_id")
    if not exam_id:
        latest = await collection(C.EXAMS).find_one(
            {"tenant_id": tenant.id, "is_deleted": {"$ne": True}}, sort=[("start_date", -1)]
        )
        if latest is None:
            raise NotFound("No exams have been created yet")
        exam_id = str(latest["_id"])

    exam = await collection(C.EXAMS).find_one({"_id": ObjectId(exam_id), "tenant_id": tenant.id})
    if exam is None:
        raise NotFound("Exam not found")

    grouped = await collection(C.MARKS).aggregate([
        {"$match": {"tenant_id": tenant.id, "exam_id": ObjectId(exam_id),
                    "is_deleted": {"$ne": True}}},
        {"$group": {
            "_id": "$subject_id",
            "students": {"$sum": 1},
            "average": {"$avg": "$marks_obtained"},
            "highest": {"$max": "$marks_obtained"},
            "lowest": {"$min": "$marks_obtained"},
            "max_marks": {"$max": "$max_marks"},
            "passed": {"$sum": {"$cond": ["$is_pass", 1, 0]}},
        }},
    ]).to_list(length=None)

    subjects = await _subject_labels(tenant.id)
    rows = []
    for group in grouped:
        maximum = float(group.get("max_marks") or 100)
        average = round(float(group.get("average") or 0), 2)
        rows.append({
            "subject": subjects.get(group["_id"], "—"),
            "students": group["students"],
            "average": average,
            "average_pct": round(average / maximum * 100, 2) if maximum else 0,
            "highest": round(float(group.get("highest") or 0), 2),
            "lowest": round(float(group.get("lowest") or 0), 2),
            "pass_rate": round(group["passed"] / group["students"] * 100, 2)
            if group["students"] else 0,
        })
    rows.sort(key=lambda r: r["average_pct"], reverse=True)

    return Report(
        key="exam_performance",
        title=f"Exam Performance — {exam.get('name', '')}",
        description="Subject averages and pass rates.",
        columns=[
            Column("subject", "Subject"),
            Column("students", "Students", "number", "right"),
            Column("average", "Average", "number", "right"),
            Column("average_pct", "Average %", "percent", "right"),
            Column("highest", "Highest", "number", "right"),
            Column("lowest", "Lowest", "number", "right"),
            Column("pass_rate", "Pass rate", "percent", "right"),
        ],
        rows=rows,
        summary=[
            {"label": "Subjects", "value": len(rows)},
            {"label": "Overall average",
             "value": round(sum(r["average_pct"] for r in rows) / max(len(rows), 1), 2),
             "type": "percent"},
            {"label": "Overall pass rate",
             "value": round(sum(r["pass_rate"] for r in rows) / max(len(rows), 1), 2),
             "type": "percent"},
        ],
        charts=[{"type": "bar", "title": "Average by subject",
                 "data": [{"name": r["subject"], "value": r["average_pct"]} for r in rows]}],
    )


async def student_marksheet(tenant: TenantContext, **params) -> Report:
    exam_id = params.get("exam_id")
    if not exam_id:
        latest = await collection(C.EXAMS).find_one(
            {"tenant_id": tenant.id, "is_deleted": {"$ne": True}}, sort=[("start_date", -1)]
        )
        if latest is None:
            raise NotFound("No exams have been created yet")
        exam_id = str(latest["_id"])

    match: dict[str, Any] = {"tenant_id": tenant.id, "exam_id": ObjectId(exam_id),
                             "is_deleted": {"$ne": True}}
    if params.get("section_id"):
        match["section_id"] = ObjectId(params["section_id"])
    elif params.get("class_id"):
        match["class_id"] = ObjectId(params["class_id"])
    else:
        # A marksheet ranks students against each other, so it has to be one
        # class. Unfiltered, pick the class with the most marks entered rather
        # than ranking Class 7 against Class 8 on different papers.
        busiest = await collection(C.MARKS).aggregate([
            {"$match": match},
            {"$group": {"_id": "$class_id", "n": {"$sum": 1}}},
            {"$sort": {"n": -1}}, {"$limit": 1},
        ]).to_list(length=1)
        if busiest and busiest[0]["_id"]:
            match["class_id"] = busiest[0]["_id"]

    marks = await collection(C.MARKS).find(match).to_list(length=20000)
    if not marks:
        return Report(key="student_marksheet", title="Class Marksheet",
                      description="No marks entered for this selection.",
                      columns=[Column("student", "Student")], rows=[])

    subjects = await _subject_labels(tenant.id)
    subject_ids = sorted({m["subject_id"] for m in marks}, key=lambda s: subjects.get(s, ""))

    students = {
        s["_id"]: s
        for s in await collection(C.STUDENTS).find(
            {"_id": {"$in": list({m["student_id"] for m in marks})}},
            {"first_name": 1, "last_name": 1, "roll_number": 1, "admission_number": 1},
        ).to_list(length=None)
    }

    per_student: dict[ObjectId, dict[str, Any]] = {}
    for mark in marks:
        entry = per_student.setdefault(mark["student_id"], {"obtained": 0.0, "max": 0.0,
                                                            "subjects": {}})
        obtained = float(mark.get("marks_obtained") or 0)
        maximum = float(mark.get("max_marks") or 0)
        entry["subjects"][str(mark["subject_id"])] = obtained
        entry["obtained"] += obtained
        entry["max"] += maximum

    rows = []
    for student_id, entry in per_student.items():
        student = students.get(student_id, {})
        percentage = round(entry["obtained"] / entry["max"] * 100, 2) if entry["max"] else 0
        row = {
            "roll_number": student.get("roll_number", ""),
            "student": " ".join(filter(None, [student.get("first_name"),
                                              student.get("last_name")])),
            "total": round(entry["obtained"], 2),
            "out_of": round(entry["max"], 2),
            "percentage": percentage,
        }
        for subject_id in subject_ids:
            row[str(subject_id)] = entry["subjects"].get(str(subject_id), "—")
        rows.append(row)

    rows.sort(key=lambda r: r["percentage"], reverse=True)
    for index, row in enumerate(rows, start=1):
        row["rank"] = index

    classes = await _lookup(C.CLASSES, tenant.id)
    scope = classes.get(match.get("class_id"), "")

    columns = [
        Column("rank", "Rank", "number", "center"),
        Column("roll_number", "Roll", "number", "center"),
        Column("student", "Student"),
        *[Column(str(s), subjects.get(s, "—"), "number", "right") for s in subject_ids],
        Column("total", "Total", "number", "right"),
        Column("percentage", "%", "percent", "right"),
    ]
    return Report(
        key="student_marksheet",
        title=f"Class Marksheet — {scope}" if scope else "Class Marksheet",
        description="Every student's marks across subjects, ranked within the class.",
        columns=columns,
        rows=rows,
        summary=[
            {"label": "Students", "value": len(rows)},
            {"label": "Class average",
             "value": round(sum(r["percentage"] for r in rows) / max(len(rows), 1), 2),
             "type": "percent"},
            {"label": "Highest", "value": rows[0]["percentage"] if rows else 0,
             "type": "percent"},
        ],
    )


async def admission_funnel(tenant: TenantContext, **params) -> Report:
    start, end = _range(params.get("start"), params.get("end"), 180)
    grouped = await collection(C.ADMISSION_ENQUIRIES).aggregate([
        {"$match": {"tenant_id": tenant.id, "is_deleted": {"$ne": True},
                    "created_at": _between(start, end)}},
        {"$group": {
            "_id": "$source",
            "total": {"$sum": 1},
            "contacted": {"$sum": {"$cond": [{"$in": ["$status",
                                                      ["contacted", "visited", "converted"]]}, 1, 0]}},
            "visited": {"$sum": {"$cond": [{"$in": ["$status", ["visited", "converted"]]}, 1, 0]}},
            "converted": {"$sum": {"$cond": [{"$eq": ["$status", "converted"]}, 1, 0]}},
            "lost": {"$sum": {"$cond": [{"$eq": ["$status", "lost"]}, 1, 0]}},
        }},
        {"$sort": {"total": -1}},
    ]).to_list(length=None)

    rows = [
        {
            "source": (group["_id"] or "unknown").replace("_", " ").title(),
            "enquiries": group["total"],
            "contacted": group["contacted"],
            "visited": group["visited"],
            "converted": group["converted"],
            "lost": group["lost"],
            "conversion": round(group["converted"] / group["total"] * 100, 2)
            if group["total"] else 0,
        }
        for group in grouped
    ]
    total = sum(r["enquiries"] for r in rows)
    converted = sum(r["converted"] for r in rows)

    return Report(
        key="admission_funnel",
        title="Admission Funnel",
        description=f"{start:%d %b %Y} to {end:%d %b %Y}",
        columns=[
            Column("source", "Source"),
            Column("enquiries", "Enquiries", "number", "right"),
            Column("contacted", "Contacted", "number", "right"),
            Column("visited", "Visited", "number", "right"),
            Column("converted", "Admitted", "number", "right"),
            Column("lost", "Lost", "number", "right"),
            Column("conversion", "Conversion", "percent", "right"),
        ],
        rows=rows,
        summary=[
            {"label": "Enquiries", "value": total},
            {"label": "Admitted", "value": converted},
            {"label": "Conversion",
             "value": round(converted / total * 100, 2) if total else 0, "type": "percent"},
        ],
        charts=[{"type": "funnel", "title": "Pipeline", "data": [
            {"name": "Enquiries", "value": total},
            {"name": "Contacted", "value": sum(r["contacted"] for r in rows)},
            {"name": "Visited", "value": sum(r["visited"] for r in rows)},
            {"name": "Admitted", "value": converted},
        ]}],
    )


async def library_circulation(tenant: TenantContext, **params) -> Report:
    start, end = _range(params.get("start"), params.get("end"), 90)
    loans = await collection(C.LIBRARY_LOANS).find({
        "tenant_id": tenant.id, "is_deleted": {"$ne": True},
        "issued_on": _between(start, end),
    }).sort([("issued_on", -1)]).to_list(length=5000)

    today = date.today()
    rows = []
    for loan in loans:
        due = loan.get("due_date")
        overdue = (
            loan.get("status") == "issued" and due is not None and due.date() < today
        )
        rows.append({
            "item": loan.get("item_title", ""),
            "borrower": loan.get("borrower_name", ""),
            "type": (loan.get("borrower_type") or "").title(),
            "issued_on": loan["issued_on"].date().isoformat() if loan.get("issued_on") else "",
            "due_date": due.date().isoformat() if due else "",
            "returned_on": loan["returned_on"].date().isoformat()
            if loan.get("returned_on") else "",
            "status": "overdue" if overdue else loan.get("status", ""),
            "fine": round(float(loan.get("fine_amount") or 0), 2),
        })

    return Report(
        key="library_circulation",
        title="Library Circulation",
        description=f"{start:%d %b %Y} to {end:%d %b %Y}",
        columns=[
            Column("item", "Item"),
            Column("borrower", "Borrower"),
            Column("type", "Type"),
            Column("issued_on", "Issued", "date"),
            Column("due_date", "Due", "date"),
            Column("returned_on", "Returned", "date"),
            Column("fine", "Fine", "money", "right"),
            Column("status", "Status", "badge", "center"),
        ],
        rows=rows,
        summary=[
            {"label": "Loans", "value": len(rows)},
            {"label": "Currently out",
             "value": sum(1 for r in rows if r["status"] in {"issued", "overdue"})},
            {"label": "Overdue", "value": sum(1 for r in rows if r["status"] == "overdue")},
            {"label": "Fines", "value": round(sum(r["fine"] for r in rows), 2), "type": "money"},
        ],
    )


async def transport_usage(tenant: TenantContext, **params) -> Report:
    routes = await collection(C.TRANSPORT_ROUTES).find(
        {"tenant_id": tenant.id, "is_deleted": {"$ne": True}}
    ).to_list(length=500)

    allocations = await collection(C.TRANSPORT_ALLOCATIONS).aggregate([
        {"$match": {"tenant_id": tenant.id, "status": "active", "is_deleted": {"$ne": True}}},
        {"$group": {"_id": "$route_id", "students": {"$sum": 1},
                    "fare": {"$sum": "$monthly_fare"}}},
    ]).to_list(length=None)
    by_route = {row["_id"]: row for row in allocations}

    vehicles = {
        v["_id"]: v
        for v in await collection(C.TRANSPORT_VEHICLES).find(
            {"tenant_id": tenant.id}
        ).to_list(length=None)
    }

    rows = []
    for route in routes:
        usage = by_route.get(route["_id"], {"students": 0, "fare": 0})
        vehicle = vehicles.get(route.get("vehicle_id"), {})
        capacity = int(vehicle.get("capacity") or 0)
        rows.append({
            "route": route.get("name", ""),
            "code": route.get("code", ""),
            "vehicle": vehicle.get("registration_number", "—"),
            "capacity": capacity or "—",
            "students": usage["students"],
            "occupancy": round(usage["students"] / capacity * 100, 2) if capacity else 0,
            "monthly_fare": round(float(usage["fare"] or 0), 2),
        })

    return Report(
        key="transport_usage",
        title="Transport Usage",
        description="Route occupancy and monthly fare value.",
        columns=[
            Column("route", "Route"),
            Column("code", "Code"),
            Column("vehicle", "Vehicle"),
            Column("capacity", "Capacity", "number", "right"),
            Column("students", "Students", "number", "right"),
            Column("occupancy", "Occupancy", "percent", "right"),
            Column("monthly_fare", "Monthly fare", "money", "right"),
        ],
        rows=rows,
        summary=[
            {"label": "Routes", "value": len(rows)},
            {"label": "Students", "value": sum(r["students"] for r in rows)},
            {"label": "Monthly value",
             "value": round(sum(r["monthly_fare"] for r in rows), 2), "type": "money"},
        ],
    )


BUILDERS = {
    "student_roster": student_roster,
    "student_demographics": student_demographics,
    "staff_directory": staff_directory,
    "attendance_register": attendance_register,
    "attendance_defaulters": attendance_defaulters,
    "daily_attendance": daily_attendance,
    "fee_collection": fee_collection,
    "fee_outstanding": fee_outstanding,
    "fee_by_class": fee_by_class,
    "expense_summary": expense_summary,
    "exam_performance": exam_performance,
    "student_marksheet": student_marksheet,
    "admission_funnel": admission_funnel,
    "library_circulation": library_circulation,
    "transport_usage": transport_usage,
}


async def build(tenant: TenantContext, key: str, **params) -> Report:
    builder = BUILDERS.get(key)
    if builder is None:
        raise NotFound(f"No report called '{key}'")
    return await builder(tenant, **params)


def to_csv(report: Report) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c.label for c in report.columns])
    for row in report.rows:
        writer.writerow([row.get(c.key, "") for c in report.columns])
    if report.summary:
        writer.writerow([])
        for item in report.summary:
            writer.writerow([item["label"], item["value"]])
    return buffer.getvalue()


def to_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.title[:31] or "Report"

    sheet.append([report.title])
    sheet["A1"].font = Font(size=14, bold=True)
    sheet.append([report.description])
    sheet.append([f"Generated {report.generated_at:%d %b %Y, %H:%M}"])
    sheet.append([])

    header_row = sheet.max_row + 1
    sheet.append([c.label for c in report.columns])
    fill = PatternFill("solid", start_color="FF111214")
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for row in report.rows:
        sheet.append([row.get(c.key, "") for c in report.columns])

    if report.summary:
        sheet.append([])
        for item in report.summary:
            sheet.append([item["label"], item["value"]])

    for index, column in enumerate(report.columns, start=1):
        longest = max(
            [len(column.label)] + [len(str(r.get(column.key, ""))) for r in report.rows[:200]]
        )
        sheet.column_dimensions[sheet.cell(row=header_row, column=index).column_letter].width = (
            min(max(longest + 3, 10), 40)
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
